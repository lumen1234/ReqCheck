"""需求验证 Word 报告导出（仅本地下载，不进共享卷）。"""
import os
import tempfile
from datetime import datetime

from flask import Blueprint, request, jsonify, send_file
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from app import app
from app.models import Document, DocumentBatch
from app.routes.export import (
    _load_requirement_tree,
    _load_validation_map,
    _ensure_classified_saved,
    _flatten_tree,
)
from app.services import project_service

report_bp = Blueprint('report', __name__)

CONTENT_MAX_LENGTH = 100
TABLE_COLS = [
    '序号', '需求标题', '需求内容', '验证结果',
    '判断依据', '测试难度', '特殊环境', 'Mock',
]


def _truncate(text: str, max_len: int = CONTENT_MAX_LENGTH) -> str:
    if not text:
        return ''
    text = text.strip()
    if len(text) <= max_len:
        return text
    return text[:max_len] + '…'


def _result_label(validation_result) -> str:
    if validation_result is True:
        return '通过'
    if validation_result is False:
        return '不通过'
    return '未验证'


def _bool_label(value) -> str:
    if value is True:
        return '是'
    if value is False:
        return '否'
    return ''


def _generate_word_report(requirements, doc_name, batch_name=None):
    """生成 Word 报告文件，返回临时文件路径。"""
    doc = DocxDocument()

    # ── 辅助：将段落所有 run 设为宋体 ──
    def _set_run_font(run, size=None):
        run.font.name = '宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        if size:
            run.font.size = size

    # ── 全局样式 ──
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    font.size = Pt(10.5)

    # ── 封面 / 标题 ──
    title = doc.add_heading('需求验证报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        _set_run_font(run)

    info_lines = [
        f'文档名称：{doc_name}',
        f'验证时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}',
        f'需求总数（含非需求节点）：{len(requirements)}',
    ]
    if batch_name:
        info_lines.insert(1, f'批次名称：{batch_name}')
    for line in info_lines:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(4)

    # ── 一、验证范围 ──
    h1 = doc.add_heading('一、验证范围', level=1)
    for run in h1.runs:
        _set_run_font(run)
    req_count = sum(1 for r in requirements if r.get('is_req') == 1)
    scope_text = (
        f'本报告对"{doc_name}"中的需求条目进行了 GJB 438C 附录 J 合规性验证。'
        f'共检查 {len(requirements)} 个节点，其中软件需求 {req_count} 条。'
    )
    doc.add_paragraph(scope_text)

    # ── 二、逐条验证结论 ──
    h2 = doc.add_heading('二、逐条验证结论', level=1)
    for run in h2.runs:
        _set_run_font(run)

    table = doc.add_table(rows=1, cols=len(TABLE_COLS), style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    hdr_cells = table.rows[0].cells
    for i, col_name in enumerate(TABLE_COLS):
        hdr_cells[i].text = col_name
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                _set_run_font(run, size=Pt(9))

    # 数据行
    for item in requirements:
        row_cells = table.add_row().cells
        is_req = item.get('is_req') == 1
        validation_result = item.get('validation_result')

        if is_req:
            result_text = _result_label(validation_result)
            reason_text = item.get('validation_reason', '')
            difficulty = item.get('test_difficulty', '')
            special_env = _bool_label(item.get('needs_special_env'))
            mock_label = _bool_label(item.get('needs_mock'))
        else:
            result_text = '非需求'
            reason_text = '此条并非需求'
            difficulty = ''
            special_env = ''
            mock_label = ''

        row_values = [
            item.get('id', ''),
            item.get('title', ''),
            _truncate(item.get('content', '')),
            result_text,
            _truncate(reason_text, 80),
            difficulty,
            special_env,
            mock_label,
        ]
        for i, val in enumerate(row_values):
            row_cells[i].text = val
            for paragraph in row_cells[i].paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, size=Pt(9))

    # ── 三、通过率统计 ──
    h3 = doc.add_heading('三、通过率统计', level=1)
    for run in h3.runs:
        _set_run_font(run)

    req_items = [r for r in requirements if r.get('is_req') == 1]
    total = len(req_items)
    passed = sum(1 for r in req_items if r.get('validation_result') is True)
    failed = sum(1 for r in req_items if r.get('validation_result') is False)
    unverified = sum(1 for r in req_items if r.get('validation_result') is None)
    pass_rate = f'{passed / total * 100:.1f}%' if total > 0 else 'N/A'

    stats = [
        f'需求总数：{total}',
        f'通过：{passed}',
        f'不通过：{failed}',
        f'未验证：{unverified}',
        f'通过率：{pass_rate}',
    ]
    for line in stats:
        doc.add_paragraph(line, style='List Bullet')

    # 按类型统计
    type_stats = {}
    for r in req_items:
        req_type = r.get('type', '未分类') or '未分类'
        type_stats[req_type] = type_stats.get(req_type, {'total': 0, 'passed': 0})
        type_stats[req_type]['total'] += 1
        if r.get('validation_result') is True:
            type_stats[req_type]['passed'] += 1

    if type_stats:
        doc.add_paragraph('')  # spacer
        p = doc.add_paragraph('按需求类型分布：')
        p.runs[0].bold = True
        for t, counts in sorted(type_stats.items()):
            t_pass_rate = f'{counts["passed"] / counts["total"] * 100:.1f}%' if counts['total'] > 0 else 'N/A'
            doc.add_paragraph(
                f'{t}：{counts["total"]} 条，通过 {counts["passed"]} 条（{t_pass_rate}）',
                style='List Bullet',
            )

    # ── 保存到临时文件 ──
    fd, tmp_path = tempfile.mkstemp(suffix='.docx', prefix='reqcheck_report_')
    os.close(fd)
    doc.save(tmp_path)
    return tmp_path


# ═══════════════════════════════════════════════════════════
#  端点
# ═══════════════════════════════════════════════════════════

@report_bp.route('/api/export/<doc_id>/word', methods=['GET'])
def export_word_report(doc_id):
    if not doc_id:
        return jsonify({'error': 'doc_id is required'}), 400

    req_tree = _load_requirement_tree(doc_id)
    if not req_tree:
        return jsonify({'error': 'Requirement tree not found'}), 404

    _ensure_classified_saved(doc_id, req_tree)
    requirements, _ = _flatten_tree(
        req_tree, _load_validation_map(doc_id), doc_number=1, counter_start=1,
    )

    # 获取文档名称
    doc_record = Document.query.filter_by(id=doc_id).first()
    doc_name = doc_record.filename if doc_record else doc_id

    tmp_path = _generate_word_report(requirements, doc_name)
    download_name = f'验证报告_{doc_name}_{datetime.now().strftime("%Y%m%d")}.docx'

    try:
        return send_file(
            tmp_path,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=download_name,
        )
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


@report_bp.route('/api/export/batch/<batch_id>/word', methods=['GET'])
def export_batch_word_report(batch_id):
    portal_project_id = request.args.get('portal_project_id') or None

    batch = DocumentBatch.query.filter_by(id=batch_id).first()
    doc_entries = []
    batch_name = None

    if batch:
        batch_name = batch.name
        for document in Document.query.filter_by(batch_id=batch_id).order_by(
            Document.batch_order.asc()
        ).all():
            doc_entries.append({
                'doc_id': document.id,
                'filename': document.filename,
                'batch_order': document.batch_order or 1,
            })
    else:
        uniportal_batch = project_service.get_uniportal_batch_detail(
            batch_id, portal_project_id=portal_project_id,
        )
        if not uniportal_batch:
            return jsonify({'error': 'Batch not found'}), 404
        batch_name = uniportal_batch.get('batch_name') or batch_id
        for doc in uniportal_batch.get('documents') or []:
            doc_entries.append({
                'doc_id': doc['doc_id'],
                'filename': doc['filename'],
                'batch_order': doc.get('doc') or 1,
            })

    if not doc_entries:
        return jsonify({'error': 'Batch has no documents'}), 404

    # 扁平化所有文档
    requirements = []
    counter = 1
    missing = []
    for document in doc_entries:
        req_tree = _load_requirement_tree(document['doc_id'])
        if not req_tree:
            missing.append(document['filename'])
            continue
        _ensure_classified_saved(document['doc_id'], req_tree)
        flattened, counter = _flatten_tree(
            req_tree,
            _load_validation_map(document['doc_id']),
            doc_number=document['batch_order'],
            counter_start=counter,
            node_id_prefix=f"{document['doc_id']}:",
        )
        for item in flattened:
            item['doc_id'] = document['doc_id']
            item['source_filename'] = document['filename']
        requirements.extend(flattened)

    if not requirements:
        return jsonify({'error': 'No parsed requirement trees found', 'missing': missing}), 404

    tmp_path = _generate_word_report(
        requirements,
        batch_name or batch_id,
        batch_name=batch_name,
    )
    download_name = f'验证报告_批次{batch_name or batch_id}_{datetime.now().strftime("%Y%m%d")}.docx'

    try:
        return send_file(
            tmp_path,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=download_name,
        )
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
